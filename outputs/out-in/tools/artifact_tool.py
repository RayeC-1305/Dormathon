"""Standalone, DEMO-only invoice-financing artifacts; no money movement.

Input contract:
* loan_terms: principal (aliases loan_amount/amount), annual_interest_rate as
  a decimal FRACTION (0.12 means 12%), term_months (alias months), and optional
  loan_cap (defaults explicitly to principal). Money accepts Decimal, integers,
  or exact decimal strings, never binary floats. Optional first_repayment_date
  is ISO YYYY-MM-DD; borrower/lender/currency and other terms are only printed
  when supplied. annual_interest_rate_pct is an explicit percent alternative.
* invoices: unique invoice_id (aliases id/invoice_number), items (aliases
  line_items/lines) containing quantity (alias qty) and unit_price. Each exact
  product must be cent-valued; sub-cent valuations are rejected, not rounded.
  Required subtotal (alias invoice_subtotal) must equal the exact line sum.
  Required positive total (aliases amount/invoice_total) must equal subtotal
  + tax + freight - discount. Adjustments are nonnegative exact cent amounts;
  omitted adjustments contribute zero under the explicit no-adjustment policy.
  Optional description and customer/borrower are supplied evidence only.
* bank_rows: invoice_id, optional date/reference/description, and inflow,
  credit, or signed amount. Negative signed amounts are outflows, not inflows.
* build_artifacts obtains bank_rows from audit['bank_rows'] if supplied.

DecimalAudit in the workbook and decimal_audit in the canonical manifest are
nonauthoritative about real-world truth, but authoritative for deterministic
arithmetic. Excel's binary-float formulas are a display/review surface, not an
accounting authority. Detached Ed25519 attestations demonstrate byte integrity,
NOT identity, legal countersignature, underwriting approval, or real funding.
Dependencies are lazy-loaded; importing this module has no side effects.
"""

from __future__ import annotations

import calendar
import hashlib
import io
import json
import math
import os
import re
import stat
import tempfile
from datetime import date
from decimal import Context, Decimal, InvalidOperation, ROUND_DOWN, ROUND_HALF_UP, localcontext
from pathlib import Path, PurePosixPath
from xml.sax.saxutils import escape

CENT = Decimal('0.01')
ZERO = Decimal('0')
PRECISION = 50
DECIMAL_CONTEXT = Context(prec=PRECISION, rounding=ROUND_HALF_UP)
LEGAL_STATUS = 'demo_attestation_not_legal_countersignature'
DEMO_NOTICE = 'DEMO - NOT LEGALLY COUNTERSIGNED - NO REAL FUNDS'
PDF_NAME = 'Out_In_TNC_Agreement.pdf'
LEDGER_NAME = 'MasterLedger.xlsx'
MANIFEST_NAME = 'verification_manifest.json'
MANIFEST_SIG = 'verification_manifest.sig'
PUBLIC_KEYS = {
    'manifest': 'demo_public_key.pem',
    'borrower': 'demo_borrower_public_key.pem',
    'lender': 'demo_lender_public_key.pem',
}
ROLE_SIGNATURES = {
    'borrower': 'borrower_attestation.sig',
    'lender': 'lender_attestation.sig',
}
ATTESTATION_DOMAIN = b'OUT_IN_DEMO_PDF_ATTESTATION_V1\x00'
POLICY = {
    'decimal_precision': PRECISION,
    'decimal_rounding': 'ROUND_HALF_UP in an isolated context, independent of caller settings',
    'amount_inputs': 'exact decimal strings, Decimal, or integers; floats rejected',
    'annual_rate_unit': 'fraction; annual_interest_rate_pct explicitly accepts percent',
    'monthly_rate': 'annual simple interest rate / Decimal(12); no compounding',
    'interest_basis': 'original principal throughout term, not declining balance',
    'total_interest': 'principal * annual_rate * term_months / 12, ROUND_HALF_UP to cents',
    'repayments': 'equal floor-to-cent principal and interest; last installment gets remainder',
    'allocations': 'principal * adjusted invoice total / adjusted invoice total sum, floor cents; last gets remainder',
    'invoice_lines': 'exact Decimal quantity * unit_price; reject sub-cent products; no line rounding',
    'invoice_subtotal': 'required declared subtotal equals exact sum of unrounded cent-valued products',
    'invoice_total': 'required positive declared total equals subtotal + tax + freight - discount',
    'invoice_adjustments': 'tax/freight/discount are nonnegative exact cents; omitted means zero contribution',
    'money_inputs': 'principal, cap, invoice subtotals/totals/adjustments, and bank amounts must be exact cents',
    'annual_rate_bounds': '0 <= annual fraction <= 1 (explicit percent alternative: 0 through 100)',
    'bank_inflow': 'max(supplied inflow/credit/signed amount, 0), exact invoice-id match',
    'loan_cap_default': 'if omitted, use supplied principal; no fabricated approval',
    'due_dates': 'omitted unless first_repayment_date supplied; anchor day clamped each month',
    'excel': 'binary floats are display only; do not round reconciliation balances',
    'authority': 'Decimal canonical JSON, not cached Excel formula results',
    'legal_status': LEGAL_STATUS,
    'real_funds': False,
}


def _json_safe(value):
    if isinstance(value, Decimal):
        if not value.is_finite():
            raise ValueError('Non-finite Decimal evidence is forbidden')
        return format(value, 'f')
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError('Non-finite evidence is forbidden')
        # Preserve supplied nonmonetary evidence without introducing JSON floats.
        return str(value)
    if isinstance(value, dict):
        if any(not isinstance(key, str) for key in value):
            raise TypeError('Evidence object keys must be strings')
        return {key: _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    raise TypeError(f'Unsupported evidence type: {type(value).__name__}')


def canonical_json(value) -> bytes:
    """Canonical UTF-8 JSON used by all evidence and manifest signatures."""
    return json.dumps(_json_safe(value), sort_keys=True, separators=(',', ':'),
                      ensure_ascii=False, allow_nan=False).encode('utf-8')


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def _hash_string(value: str) -> str:
    if not isinstance(value, str) or not re.fullmatch(r'[0-9a-fA-F]{64}', value):
        raise ValueError('SHA256 must be a 64-character hexadecimal string')
    return value.lower()


def _relative_name(name: str) -> str:
    if not isinstance(name, str) or not name or '\\' in name or ':' in name:
        raise ValueError('Evidence names must be safe relative POSIX paths')
    path = PurePosixPath(name)
    if path.is_absolute() or any(part in ('', '.', '..') for part in name.split('/')):
        raise ValueError('Evidence names cannot be absolute or contain traversal')
    return name


def _pick(mapping: dict, names: tuple, *, default=None, required=False):
    supplied = [(key, mapping[key]) for key in names if key in mapping and mapping[key] is not None]
    if not supplied:
        if required:
            raise ValueError(f'Missing required field: {names[0]}')
        return default
    if any(value != supplied[0][1] for _, value in supplied[1:]):
        raise ValueError(f'Conflicting aliases for {names[0]}')
    return supplied[0][1]


def _decimal(value, field: str, *, cents=False, nonnegative=True) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (str, int, Decimal)):
        raise TypeError(f'{field} must be an exact decimal string, Decimal, or integer')
    if isinstance(value, str) and not re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)', value):
        raise ValueError(f'{field} must be a plain decimal string')
    try:
        result = Decimal(value)
    except (ValueError, InvalidOperation) as exc:
        raise ValueError(f'Invalid decimal for {field}') from exc
    if not result.is_finite() or (nonnegative and result < 0):
        raise ValueError(f'{field} must be finite and nonnegative')
    # Bound arithmetic and Excel display inputs; never silently lose precision.
    if len(result.as_tuple().digits) > 28 or result.adjusted() > 15 or result.as_tuple().exponent < -28:
        raise ValueError(f'{field} exceeds supported Decimal precision/range')
    if cents:
        with localcontext(DECIMAL_CONTEXT) as ctx:
            ctx.prec = PRECISION
            if result != result.quantize(CENT):
                raise ValueError(f'{field} must be an exact cent amount')
    return result.copy_abs() if result == 0 else result


def _money(value: Decimal) -> str:
    return format(value, '.2f')


def _identifier(value, field: str) -> str:
    if not isinstance(value, (str, int)) or isinstance(value, bool) or not str(value):
        raise ValueError(f'{field} must be a nonempty string or integer')
    return str(value)


def _parts(total: Decimal, weights: list[Decimal]) -> list[Decimal]:
    """Floor proportional parts; preserve every cent in the last remainder."""
    if not weights:
        if total:
            raise ValueError('Cannot allocate a nonzero amount without weights')
        return []
    denominator = sum(weights, ZERO)
    if not denominator:
        if total:
            raise ValueError('Cannot allocate a nonzero amount to zero weights')
        return [ZERO.quantize(CENT) for _ in weights]
    result = [(total * weight / denominator).quantize(CENT, rounding=ROUND_DOWN)
              for weight in weights[:-1]]
    return result + [total - sum(result, ZERO)]


def _loan_model(loan_terms: dict) -> dict:
    if not isinstance(loan_terms, dict):
        raise TypeError('loan_terms must be a dictionary')
    with localcontext(DECIMAL_CONTEXT) as ctx:
        ctx.prec = PRECISION
        principal = _decimal(_pick(loan_terms, ('principal', 'loan_amount', 'amount'), required=True),
                             'principal', cents=True)
        cap = _decimal(_pick(loan_terms, ('loan_cap', 'cap'), default=principal), 'loan_cap', cents=True)
        if principal > cap:
            raise ValueError('Demo principal exceeds supplied loan cap')
        annual_value = _pick(loan_terms, ('annual_interest_rate', 'annual_rate'))
        percent_value = loan_terms.get('annual_interest_rate_pct')
        if annual_value is None and percent_value is None:
            raise ValueError('Missing annual_interest_rate (fraction) or annual_interest_rate_pct')
        annual = (_decimal(annual_value, 'annual_interest_rate') if annual_value is not None
                  else _decimal(percent_value, 'annual_interest_rate_pct') / Decimal(100))
        if percent_value is not None and annual != _decimal(percent_value, 'annual_interest_rate_pct') / Decimal(100):
            raise ValueError('Conflicting fraction and percent annual interest rates')
        if annual > Decimal(1):
            raise ValueError('annual_interest_rate must be a fraction from 0 through 1 (0% through 100%)')
        months_decimal = _decimal(_pick(loan_terms, ('term_months', 'months'), required=True), 'term_months')
        if months_decimal != months_decimal.to_integral_value() or not 1 <= months_decimal <= 1200:
            raise ValueError('term_months must be an integer from 1 to 1200')
        months = int(months_decimal)
        if len(principal.as_tuple().digits) + len(annual.as_tuple().digits) + len(str(months)) > PRECISION - 2:
            raise ValueError('Loan interest arithmetic exceeds supported precision')
        monthly = annual / Decimal(12)
        interest = (principal * annual * Decimal(months) / Decimal(12)).quantize(CENT, rounding=ROUND_HALF_UP)
        principal_parts = _parts(principal, [Decimal(1)] * months)
        interest_parts = _parts(interest, [Decimal(1)] * months)
        first = loan_terms.get('first_repayment_date')
        if first is not None:
            if not isinstance(first, str) or not re.fullmatch(r'\d{4}-\d{2}-\d{2}', first):
                raise ValueError('first_repayment_date must be ISO YYYY-MM-DD')
            first = date.fromisoformat(first)
        balance = principal
        schedule = []
        for index, (pmt, charge) in enumerate(zip(principal_parts, interest_parts)):
            entry = {'period': index + 1, 'opening_balance': _money(balance),
                     'principal': _money(pmt), 'interest': _money(charge),
                     'payment': _money(pmt + charge), 'closing_balance': _money(balance - pmt)}
            if first is not None:
                ordinal = first.year * 12 + first.month - 1 + index
                year, month_index = divmod(ordinal, 12)
                month = month_index + 1
                entry['due_date'] = date(year, month, min(first.day, calendar.monthrange(year, month)[1])).isoformat()
            schedule.append(entry)
            balance -= pmt
        return {'principal': _money(principal), 'loan_cap': _money(cap),
                'annual_interest_rate': format(annual, 'f'), 'monthly_interest_rate': format(monthly, 'f'),
                'term_months': months, 'total_interest': _money(interest),
                'total_repayment': _money(principal + interest), 'repayments': schedule,
                'supplied_terms': _json_safe(loan_terms), 'policy': dict(POLICY)}


def _ledger_model(invoice_data: list, bank_rows: list, loan_terms: dict) -> dict:
    if not isinstance(invoice_data, list) or not isinstance(bank_rows, list):
        raise TypeError('invoice_data and bank_rows must be lists')
    loan = _loan_model(loan_terms)
    with localcontext(DECIMAL_CONTEXT) as ctx:
        ctx.prec = PRECISION
        invoices, seen = [], set()
        for invoice in invoice_data:
            if not isinstance(invoice, dict):
                raise TypeError('Each invoice must be a dictionary')
            invoice_id = _identifier(_pick(invoice, ('invoice_id', 'id', 'invoice_number'), required=True), 'invoice_id')
            if invoice_id in seen:
                raise ValueError(f'Duplicate invoice_id: {invoice_id}')
            seen.add(invoice_id)
            items = _pick(invoice, ('items', 'line_items', 'lines'), required=True)
            if not isinstance(items, list) or not items:
                raise ValueError(f'Invoice {invoice_id} needs supplied line items')
            lines, products = [], []
            for item in items:
                if not isinstance(item, dict):
                    raise TypeError('Each invoice item must be a dictionary')
                qty = _decimal(_pick(item, ('quantity', 'qty'), required=True), 'quantity')
                price = _decimal(_pick(item, ('unit_price',), required=True), 'unit_price')
                if len(qty.as_tuple().digits) + len(price.as_tuple().digits) > PRECISION - 2:
                    raise ValueError('Line arithmetic exceeds supported precision')
                product = qty * price
                if product % CENT != ZERO:
                    raise ValueError(f'Invoice {invoice_id} has a sub-cent line valuation; exact cents required')
                products.append(product)
                lines.append({'description': item.get('description', ''), 'quantity': format(qty, 'f'),
                              'unit_price': format(price, 'f'), 'amount': _money(product)})
            subtotal = sum(products, ZERO)
            declared_subtotal = _decimal(_pick(invoice, ('subtotal', 'invoice_subtotal'), required=True),
                                         'declared invoice subtotal', cents=True)
            if declared_subtotal != subtotal:
                raise ValueError(f'Invoice {invoice_id} declared subtotal does not equal exact line sum')
            adjustments = {key: _decimal(_pick(invoice, (key,), default=ZERO),
                                         f'invoice {key}', cents=True)
                           for key in ('tax', 'freight', 'discount')}
            total = subtotal + adjustments['tax'] + adjustments['freight'] - adjustments['discount']
            declared_total = _decimal(_pick(invoice, ('total', 'amount', 'invoice_total'), required=True),
                                      'declared invoice total', cents=True)
            if declared_total <= ZERO:
                raise ValueError(f'Invoice {invoice_id} declared total must be positive')
            if total != declared_total:
                raise ValueError(f'Invoice {invoice_id} declared total does not equal subtotal + tax + freight - discount')
            invoices.append({'invoice_id': invoice_id,
                             'customer': _pick(invoice, ('customer', 'borrower'), default=''),
                             'lines': lines, 'subtotal': _money(subtotal),
                             **{key: _money(value) for key, value in adjustments.items()},
                             'total': _money(total), 'declared_subtotal': _money(declared_subtotal),
                             'declared_subtotal_difference': _money(subtotal - declared_subtotal),
                             'declared_total': _money(declared_total),
                             'declared_difference': _money(total - declared_total),
                             'matched_inflow': '0.00'})
        banks, inflows = [], {entry['invoice_id']: ZERO for entry in invoices}
        for row in bank_rows:
            if not isinstance(row, dict):
                raise TypeError('Each bank row must be a dictionary')
            # These alternatives have different semantics; explicit inflow wins.
            field = next((key for key in ('inflow', 'credit', 'amount') if row.get(key) is not None), None)
            if field is None:
                raise ValueError('Bank row needs inflow, credit, or signed amount')
            amount = _decimal(row[field], f'bank {field}', cents=True, nonnegative=field != 'amount')
            inflow = max(amount, ZERO)
            invoice_id = row.get('invoice_id')
            invoice_id = '' if invoice_id is None or invoice_id == '' else _identifier(invoice_id, 'bank invoice_id')
            matched = invoice_id in inflows
            if matched:
                inflows[invoice_id] += inflow
            banks.append({'date': row.get('date', ''), 'reference': row.get('reference', ''),
                          'description': row.get('description', ''), 'invoice_id': invoice_id,
                          'amount_field': field, 'supplied_amount': _money(amount),
                          'inflow': _money(inflow), 'matched': matched})
        support = sum((Decimal(entry['total']) for entry in invoices), ZERO)
        principal = Decimal(loan['principal'])
        if principal > support:
            raise ValueError('Demo principal exceeds computed adjusted invoice total support')
        allocations = _parts(principal, [Decimal(entry['total']) for entry in invoices])
        for entry, allocation in zip(invoices, allocations):
            entry['allocation'] = _money(allocation)
            entry['matched_inflow'] = _money(inflows[entry['invoice_id']])
            entry['inflow_difference'] = _money(inflows[entry['invoice_id']] - Decimal(entry['total']))
        return {'loan': loan, 'invoices': _json_safe(invoices), 'bank_rows': _json_safe(banks),
                'invoice_sum': _money(support), 'allocation_sum': _money(sum(allocations, ZERO)),
                'unmatched_bank_inflow': _money(sum((Decimal(row['inflow']) for row in banks if not row['matched']), ZERO)),
                'legal_status': LEGAL_STATUS, 'real_funds': False}


def _pdf_dependencies():
    try:
        from pypdf import PdfReader, PdfWriter
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise RuntimeError('PDF generation requires installed pypdf and reportlab; no packages are installed automatically') from exc
    return PdfReader, PdfWriter, colors, A4, getSampleStyleSheet, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, canvas


def compile_tnc_agreement(loan_terms: dict, template_path: str | None,
                          output_path: str, evidence_hash: str) -> dict:
    """Append a demo terms sheet to a template (or create a demo-only contract).

    evidence_hash must hash upstream sources/terms, NEVER this final PDF.
    Returns the final PDF SHA256 and unsigned demo-attestation status; signing
    is performed only after all PDF bytes are final, by build_artifacts.
    """
    loan = _loan_model(loan_terms)
    evidence_hash = _hash_string(evidence_hash)
    output = Path(output_path)
    if template_path is not None and Path(template_path).resolve() == output.resolve():
        raise ValueError('Output cannot overwrite the supplied template')
    PdfReader, PdfWriter, colors, A4, styles_factory, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, canvas = _pdf_dependencies()
    styles = styles_factory()
    story = [Paragraph(DEMO_NOTICE, styles['Title']), Spacer(1, 10),
             Paragraph('Demo terms sheet' if template_path else 'Demonstration financing contract', styles['Heading1']),
             Paragraph('This document demonstrates a proposed financing workflow only. No real funds '
                       'are transferred or committed. Any detached signatures are demo attestations of '
                       'bytes, not borrower or lender identity, legal countersignatures, or approval. '
                       'A supplied template remains reference material; this demo does not activate it.', styles['BodyText']),
             Spacer(1, 12)]

    def paragraph(value):
        return Paragraph(escape(str(value)), styles['BodyText'])

    terms = [('Principal', loan['principal']), ('Demo loan cap', loan['loan_cap']),
             ('Annual simple interest rate (fraction)', loan['annual_interest_rate']),
             ('Monthly rate = annual / 12 (Decimal precision 50)', loan['monthly_interest_rate']),
             ('Term in months', loan['term_months']), ('Total simple interest', loan['total_interest']),
             ('Total demo repayment', loan['total_repayment'])]
    table = Table([[paragraph(label), paragraph(value)] for label, value in terms], colWidths=[255, 255])
    table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
                              ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                              ('LEFTPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8)]))
    story.extend([table, Spacer(1, 14), Paragraph('Only supplied additional terms', styles['Heading2'])])
    core_fields = {'principal', 'loan_amount', 'amount', 'loan_cap', 'cap', 'annual_interest_rate',
                   'annual_rate', 'annual_interest_rate_pct', 'term_months', 'months'}
    for key, value in sorted(loan['supplied_terms'].items()):
        if key not in core_fields and value is not None and value != '':
            text = canonical_json(value).decode('utf-8') if isinstance(value, (dict, list)) else str(value)
            story.append(paragraph(f'{key}: {text}'))
    story.extend([Spacer(1, 12), Paragraph('Explicit arithmetic policy', styles['Heading2']),
                  paragraph('Simple interest uses original principal for every month, annual fraction / 12, '
                            'without compounding. Total term interest is rounded HALF_UP to cents once. '
                            'Principal and interest installments are floored to cents; the last installment '
                            'takes each remainder. Monthly rate display uses Decimal precision 50. '
                            'Dates and borrower terms are omitted when not supplied.'),
                  Spacer(1, 8), paragraph(f'Upstream source/terms evidence SHA256: {evidence_hash}'),
                  paragraph('This is NOT a hash of this PDF. The final verification manifest hashes the completed PDF.'),
                  Spacer(1, 14), Paragraph('Demo repayments (exact decimal amounts)', styles['Heading2'])])
    repayment_rows = [['Period', 'Principal', 'Interest', 'Payment', 'Closing', 'Due date']]
    repayment_rows += [[str(row['period']), row['principal'], row['interest'], row['payment'],
                        row['closing_balance'], row.get('due_date', '')] for row in loan['repayments']]
    repayments = Table(repayment_rows, repeatRows=1, colWidths=[42, 92, 86, 92, 100, 98])
    repayments.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#172B4D')),
                                   ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                                   ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
                                   ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                                   ('FONTSIZE', (0, 0), (-1, -1), 8), ('ALIGN', (1, 1), (4, -1), 'RIGHT'),
                                   ('BOTTOMPADDING', (0, 0), (-1, -1), 7)]))
    story.append(repayments)
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=42, rightMargin=42,
                      topMargin=52, bottomMargin=52, title=DEMO_NOTICE).build(story)
    writer = PdfWriter()
    if template_path is not None:
        template = PdfReader(str(template_path))
        if template.is_encrypted:
            raise ValueError('Encrypted templates are not supported')
        for page in template.pages:
            writer.add_page(page)
    for page in PdfReader(io.BytesIO(buffer.getvalue())).pages:
        writer.add_page(page)
    # Mark EVERY page, including template pages, so reference terms cannot be
    # detached and mistaken for an executed agreement.
    for page in writer.pages:
        mark = io.BytesIO()
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        stamp = canvas.Canvas(mark, pagesize=(width, height))
        stamp.setFillColor(colors.HexColor('#B91C1C'))
        stamp.setFont('Helvetica-Bold', 10)
        stamp.drawCentredString(width / 2, height - 23, DEMO_NOTICE)
        stamp.setFont('Helvetica-Bold', 8)
        stamp.drawCentredString(width / 2, 22, 'DEMO ATTESTATIONS ONLY - NO IDENTITY OR LEGAL SIGNATURE CLAIM')
        stamp.save()
        page.merge_page(PdfReader(io.BytesIO(mark.getvalue())).pages[0])
    writer.add_metadata({'/Title': DEMO_NOTICE, '/Subject': LEGAL_STATUS,
                         '/EvidenceSHA256': evidence_hash})
    with output.open('wb') as stream:
        writer.write(stream)
    return {'name': output.name, 'sha256': _sha256(output), 'evidence_hash': evidence_hash,
            'legal_status': LEGAL_STATUS, 'sign_status': 'unsigned_pending_demo_attestation',
            'real_funds': False, 'loan': loan}


def build_master_ledger(invoice_data: list, bank_rows: list, loan_terms: dict,
                        output_path: str) -> None:
    """Build linked Excel review formulas and embedded canonical Decimal audit."""
    model = _ledger_model(invoice_data, bank_rows, loan_terms)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.workbook.properties import CalcProperties
    except ImportError as exc:
        raise RuntimeError('Workbook generation requires installed openpyxl; no packages are installed automatically') from exc
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')
    sheets = {name: workbook.create_sheet(name) for name in
              ('Loan', 'InvoiceLines', 'Invoices', 'Bank', 'Repayments', 'Controls', 'DecimalAudit')}

    def text(cell, value):
        if value is None:
            value = ''
        if not isinstance(value, str):
            value = canonical_json(value).decode('utf-8')
        if len(value) > 32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value):
            raise ValueError('Spreadsheet text contains unsupported characters or exceeds cell length')
        cell.value = value
        # Explicit text type prevents all user strings (including =, +, -, @)
        # from becoming formulas, without changing invoice IDs used for matching.
        cell.data_type = 's'

    def number(cell, value, money=True):
        # Excel numbers are deliberately a display approximation; exact strings
        # remain alongside them and in DecimalAudit / the signed manifest.
        decimal = Decimal(str(value))
        display = float(decimal)
        if not math.isfinite(display):
            raise ValueError('Amount cannot be displayed as a finite Excel number')
        cell.value = display
        cell.number_format = '#,##0.00;[Red](#,##0.00)' if money else '0.################'

    def heading(sheet, values):
        sheet.append(values)

    loan = model['loan']
    sheet = sheets['Loan']
    heading(sheet, ['Parameter', 'Linked Excel value (DEMO)', 'Exact Decimal / policy'])
    rows = [('Status', DEMO_NOTICE, DEMO_NOTICE),
            ('Principal', loan['principal'], loan['principal']),
            ('Loan cap', loan['loan_cap'], loan['loan_cap']),
            ('Invoice support', None, model['invoice_sum']),
            ('Simulated disbursement - NO REAL FUNDS', None, loan['principal']),
            ('Annual simple interest fraction', loan['annual_interest_rate'], loan['annual_interest_rate']),
            ('Monthly fraction = annual / 12', None, loan['monthly_interest_rate']),
            ('Term months', loan['term_months'], str(loan['term_months'])),
            ('Total term interest', None, loan['total_interest']),
            ('Total repayment', None, loan['total_repayment'])]
    for row_index, (label, value, exact) in enumerate(rows, 2):
        text(sheet.cell(row_index, 1), label)
        text(sheet.cell(row_index, 3), exact)
        if value is not None:
            if row_index == 2:
                text(sheet.cell(row_index, 2), value)
            else:
                number(sheet.cell(row_index, 2), value, money=row_index not in (7, 8, 9))
    invoice_end = max(2, len(model['invoices']) + 1)
    sheet['B5'] = f'=SUM(Invoices!H2:H{invoice_end})'
    sheet['B6'] = '=MIN(B3,B4,B5)'
    sheet['B8'] = '=B7/12'
    sheet['B8'].number_format = '0.################'
    sheet['B10'] = '=ROUND(B6*B7*B9/12,2)'
    sheet['B11'] = '=B6+B10'
    for row_index in (5, 6, 10, 11):
        sheet.cell(row_index, 2).number_format = '#,##0.00'
    text(sheet.cell(12, 1), 'Arithmetic authority')
    text(sheet.cell(12, 2), 'Excel floats are display only. Canonical Decimal JSON is authoritative.')
    text(sheet.cell(13, 1), 'No rounded reconciliation balances')
    text(sheet.cell(13, 2), 'Any binary-float residual remains visible in Controls.')
    for row_index, (key, value) in enumerate(sorted(POLICY.items()), 14):
        text(sheet.cell(row_index, 1), key)
        text(sheet.cell(row_index, 3), value)
    start = sheet.max_row + 1
    for row_index, (key, value) in enumerate(sorted(loan['supplied_terms'].items()), start):
        text(sheet.cell(row_index, 1), f'Supplied term: {key}')
        text(sheet.cell(row_index, 3), value)

    lines_sheet = sheets['InvoiceLines']
    heading(lines_sheet, ['Invoice ID', 'Line', 'Description', 'Quantity', 'Unit price',
                          'Line amount', 'Exact Decimal amount', 'Unrounded difference', 'Exact invoice ID match index'])
    line_row = 2
    for invoice in model['invoices']:
        for index, line in enumerate(invoice['lines'], 1):
            text(lines_sheet.cell(line_row, 1), invoice['invoice_id'])
            lines_sheet.cell(line_row, 2, index)
            text(lines_sheet.cell(line_row, 3), line['description'])
            number(lines_sheet.cell(line_row, 4), line['quantity'], money=False)
            number(lines_sheet.cell(line_row, 5), line['unit_price'], money=False)
            lines_sheet.cell(line_row, 6, f'=D{line_row}*E{line_row}').number_format = '#,##0.00'
            text(lines_sheet.cell(line_row, 7), line['amount'])
            lines_sheet.cell(line_row, 8, f'=F{line_row}-VALUE(G{line_row})')
            lines_sheet.cell(line_row, 9, f'=SUMPRODUCT(--EXACT(A{line_row},Invoices!$A$2:$A${invoice_end}),(ROW(Invoices!$A$2:$A${invoice_end})-1))')
            line_row += 1
    line_end = max(2, line_row - 1)
    bank_sheet = sheets['Bank']
    heading(bank_sheet, ['Date', 'Reference', 'Description', 'Invoice ID', 'Bank inflow',
                         'Supplied signed amount', 'Exact Decimal inflow', 'Unrounded difference', 'Source field',
                         'Exact invoice ID match index'])
    for row_index, bank in enumerate(model['bank_rows'], 2):
        for column, key in enumerate(('date', 'reference', 'description', 'invoice_id'), 1):
            text(bank_sheet.cell(row_index, column), bank[key])
        number(bank_sheet.cell(row_index, 6), bank['supplied_amount'])
        bank_sheet.cell(row_index, 5, f'=MAX(F{row_index},0)').number_format = '#,##0.00'
        text(bank_sheet.cell(row_index, 7), bank['inflow'])
        bank_sheet.cell(row_index, 8, f'=E{row_index}-VALUE(G{row_index})')
        text(bank_sheet.cell(row_index, 9), bank['amount_field'])
        bank_sheet.cell(row_index, 10, f'=SUMPRODUCT(--EXACT(D{row_index},Invoices!$A$2:$A${invoice_end}),(ROW(Invoices!$A$2:$A${invoice_end})-1))')
    bank_end = max(2, len(model['bank_rows']) + 1)

    invoice_sheet = sheets['Invoices']
    heading(invoice_sheet, ['Invoice ID', 'Supplied customer', 'Declared subtotal', 'SUMIF line subtotal',
                            'Tax', 'Freight', 'Discount', 'Adjusted total', 'Declared total',
                            'Matched bank inflow', 'Inflow - total', 'Exact Decimal subtotal',
                            'Unrounded subtotal difference', 'Exact Decimal total', 'Unrounded total difference',
                            'Demo allocation', 'Exact Decimal allocation', 'Unrounded allocation difference',
                            'Declared subtotal difference', 'Declared total difference'])
    for row_index, invoice in enumerate(model['invoices'], 2):
        text(invoice_sheet.cell(row_index, 1), invoice['invoice_id'])
        text(invoice_sheet.cell(row_index, 2), invoice['customer'])
        number(invoice_sheet.cell(row_index, 3), invoice['declared_subtotal'])
        for column, key in enumerate(('tax', 'freight', 'discount'), 5):
            number(invoice_sheet.cell(row_index, column), invoice[key])
        number(invoice_sheet.cell(row_index, 9), invoice['declared_total'])
        # SUMIF uses an EXACT-derived numeric join index: unlike direct string
        # SUMIF criteria, this distinguishes case, wildcards, numeric-looking
        # IDs, and comparison operators, consistently with the Decimal audit.
        invoice_sheet.cell(row_index, 4, f'=SUMIF(InvoiceLines!$I$2:$I${line_end},ROW()-1,InvoiceLines!$F$2:$F${line_end})')
        invoice_sheet.cell(row_index, 8, f'=D{row_index}+E{row_index}+F{row_index}-G{row_index}')
        invoice_sheet.cell(row_index, 10, f'=SUMIF(Bank!$J$2:$J${bank_end},ROW()-1,Bank!$E$2:$E${bank_end})')
        invoice_sheet.cell(row_index, 11, f'=J{row_index}-H{row_index}')
        text(invoice_sheet.cell(row_index, 12), invoice['subtotal'])
        invoice_sheet.cell(row_index, 13, f'=D{row_index}-VALUE(L{row_index})')
        text(invoice_sheet.cell(row_index, 14), invoice['total'])
        invoice_sheet.cell(row_index, 15, f'=H{row_index}-VALUE(N{row_index})')
        if row_index == invoice_end:
            allocated = '0' if row_index == 2 else f'SUM(P2:P{row_index - 1})'
            formula = f'=Loan!$B$6-{allocated}'
        else:
            formula = f'=IF(Loan!$B$5=0,0,ROUNDDOWN(Loan!$B$6*H{row_index}/Loan!$B$5,2))'
        invoice_sheet.cell(row_index, 16, formula)
        text(invoice_sheet.cell(row_index, 17), invoice['allocation'])
        invoice_sheet.cell(row_index, 18, f'=P{row_index}-VALUE(Q{row_index})')
        invoice_sheet.cell(row_index, 19, f'=D{row_index}-C{row_index}')
        invoice_sheet.cell(row_index, 20, f'=H{row_index}-I{row_index}')
        for column in (4, 8, 10, 16):
            invoice_sheet.cell(row_index, column).number_format = '#,##0.00'

    repayments = sheets['Repayments']
    heading(repayments, ['Period', 'Opening balance', 'Principal', 'Simple interest', 'Payment',
                         'Closing balance', 'Exact principal', 'Exact interest', 'Exact payment',
                         'Exact closing', 'Unrounded closing difference', 'Unrounded monthly interest', 'Due date'])
    repay_end = loan['term_months'] + 1
    for row_index, repayment in enumerate(loan['repayments'], 2):
        repayments.cell(row_index, 1, repayment['period'])
        repayments.cell(row_index, 2, '=Loan!$B$6' if row_index == 2 else f'=F{row_index - 1}')
        if row_index == repay_end:
            previous_principal = '0' if row_index == 2 else f'SUM(C2:C{row_index - 1})'
            previous_interest = '0' if row_index == 2 else f'SUM(D2:D{row_index - 1})'
            repayments.cell(row_index, 3, f'=Loan!$B$6-{previous_principal}')
            repayments.cell(row_index, 4, f'=Loan!$B$10-{previous_interest}')
        else:
            repayments.cell(row_index, 3, '=ROUNDDOWN(Loan!$B$6/Loan!$B$9,2)')
            repayments.cell(row_index, 4, '=ROUNDDOWN(Loan!$B$10/Loan!$B$9,2)')
        repayments.cell(row_index, 5, f'=C{row_index}+D{row_index}')
        repayments.cell(row_index, 6, f'=B{row_index}-C{row_index}')
        for column, key in enumerate(('principal', 'interest', 'payment', 'closing_balance'), 7):
            text(repayments.cell(row_index, column), repayment[key])
        repayments.cell(row_index, 11, f'=F{row_index}-VALUE(J{row_index})')
        repayments.cell(row_index, 12, '=Loan!$B$6*Loan!$B$8')
        text(repayments.cell(row_index, 13), repayment.get('due_date', ''))
        for column in range(2, 7):
            repayments.cell(row_index, column).number_format = '#,##0.00'

    controls = sheets['Controls']
    heading(controls, ['Control (zero difference required)', 'Unrounded formula difference',
                       'Expected exact difference', 'Status', 'Meaning'])
    checks = [
        ('Invoice total arithmetic', f'=SUM(Invoices!H2:H{invoice_end})-VALUE(Loan!C5)', 'Adjusted Decimal total reconciliation'),
        ('Disbursement matches principal', '=Loan!B6-VALUE(Loan!C3)', 'Simulated amount only'),
        ('Cap breach', '=MAX(Loan!B6-Loan!B4,0)', 'Supplied cap, not approval'),
        ('Allocation conservation', f'=SUM(Invoices!P2:P{invoice_end})-Loan!B6', 'No cent lost in adjusted-total allocations'),
        ('Principal conservation', f'=SUM(Repayments!C2:C{repay_end})-Loan!B6', 'Principal must reconcile'),
        ('Interest conservation', f'=SUM(Repayments!D2:D{repay_end})-Loan!B10', 'Rounded term interest conserved'),
        ('Repayment conservation', f'=SUM(Repayments!E2:E{repay_end})-Loan!B11', 'Total repayment reconciles'),
        ('Final outstanding balance', f'=Repayments!F{repay_end}', 'Never rounded to conceal a residual'),
        ('Bank less invoice total support', f'=SUM(Invoices!J2:J{invoice_end})-Loan!B5', 'Evidence discrepancy, not manufactured match'),
        ('Unmatched bank inflows', f'=SUM(Bank!E2:E{bank_end})-SUM(Invoices!J2:J{invoice_end})', 'Unmatched evidence remains visible'),
        ('Declared subtotal differences', f'=SUMPRODUCT(ABS(Invoices!S2:S{invoice_end}))', 'Required subtotal equals exact line sum'),
        ('Declared total differences', f'=SUMPRODUCT(ABS(Invoices!T2:T{invoice_end}))', 'Required positive total includes all adjustments'),
        ('Line Decimal discrepancies', f'=SUMPRODUCT(ABS(InvoiceLines!H2:H{line_end}))', 'Unrounded Excel product versus exact cent-valued Decimal'),
        ('Subtotal Decimal discrepancies', f'=SUMPRODUCT(ABS(Invoices!M2:M{invoice_end}))', 'Unrounded Excel subtotal versus Decimal'),
        ('Total Decimal discrepancies', f'=SUMPRODUCT(ABS(Invoices!O2:O{invoice_end}))', 'Unrounded adjusted Excel total versus Decimal'),
        ('Allocation Decimal discrepancies', f'=SUMPRODUCT(ABS(Invoices!R2:R{invoice_end}))', 'Excel display versus Decimal'),
        ('Closing Decimal discrepancies', f'=SUMPRODUCT(ABS(Repayments!K2:K{repay_end}))', 'Excel display versus Decimal'),
        ('Interest Decimal discrepancy', '=Loan!B10-VALUE(Loan!C10)', 'Excel total versus Decimal total'),
    ]
    for row_index, (label, formula, meaning) in enumerate(checks, 2):
        text(controls.cell(row_index, 1), label)
        controls.cell(row_index, 2, formula).number_format = '0.000000000000000E+00'
        controls.cell(row_index, 3, 0)
        controls.cell(row_index, 4, f'=IF(B{row_index}=C{row_index},"PASS","MISMATCH")')
        text(controls.cell(row_index, 5), meaning)
    audit_sheet = sheets['DecimalAudit']
    heading(audit_sheet, ['Part', 'Canonical JSON (join parts in order)', 'SHA256 / status'])
    audit_bytes = canonical_json(model)
    audit_text = audit_bytes.decode('utf-8')
    for row_index, offset in enumerate(range(0, len(audit_text), 30000), 2):
        audit_sheet.cell(row_index, 1, row_index - 1)
        text(audit_sheet.cell(row_index, 2), audit_text[offset:offset + 30000])
    text(audit_sheet.cell(2, 3), hashlib.sha256(audit_bytes).hexdigest())
    text(audit_sheet.cell(3, 3), DEMO_NOTICE)

    for sheet in sheets.values():
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.print_title_rows = '1:1'
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='172B4D')
            cell.font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.fill = PatternFill('solid', fgColor='F1F5F9' if cell.row % 2 == 0 else 'FFFFFF')
                cell.font = Font(name='Calibri', color='172B4D', size=11)
                cell.alignment = Alignment(vertical='top', wrap_text=cell.data_type != 'f')
        for column in range(1, sheet.max_column + 1):
            lengths = [len(str(sheet.cell(row, column).value or ''))
                       for row in range(1, sheet.max_row + 1)
                       if sheet.cell(row, column).data_type != 'f']
            sheet.column_dimensions[get_column_letter(column)].width = min(60, max(14, max(lengths, default=0) + 2))
    workbook.save(output_path)


def _crypto_dependencies():
    try:
        from cryptography.hazmat.primitives import serialization
        from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
    except ImportError as exc:
        raise RuntimeError('Demo signing requires installed cryptography; no packages are installed automatically') from exc
    return serialization, Ed25519PrivateKey


def _signing_keys(key_dir: Path) -> dict:
    """Persistent separate-role keys; private material never enters artifacts."""
    serialization, Ed25519PrivateKey = _crypto_dependencies()
    if key_dir.is_symlink():
        raise ValueError('Signing key directory must not be a symlink')
    key_dir.mkdir(mode=0o700, parents=True, exist_ok=True)
    if os.name != 'nt':
        key_dir.chmod(0o700)
    keys = {}
    for role in PUBLIC_KEYS:
        path = key_dir / f'out_in_demo_{role}_private_key.pem'
        if path.is_symlink():
            raise ValueError('Private key files must not be symlinks')
        if not path.exists():
            key = Ed25519PrivateKey.generate()
            encoded = key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                                        serialization.NoEncryption())
            descriptor, temporary = tempfile.mkstemp(prefix=f'.{role}-', suffix='.pem', dir=key_dir)
            temporary = Path(temporary)
            try:
                if os.name != 'nt':
                    os.fchmod(descriptor, 0o600)
                with os.fdopen(descriptor, 'wb') as stream:
                    stream.write(encoded)
                    stream.flush()
                    os.fsync(stream.fileno())
                # Hard-link publication is atomic and never overwrites an
                # existing key if two jobs initialize the same role at once.
                try:
                    os.link(temporary, path)
                except FileExistsError:
                    pass
            finally:
                temporary.unlink(missing_ok=True)  # Only our own temporary key.
        if not path.is_file() or path.is_symlink():
            raise ValueError('Invalid persistent private key path')
        if os.name != 'nt':
            path.chmod(0o600)
            if stat.S_IMODE(path.stat().st_mode) & 0o077:
                raise PermissionError('Private key must not be group/world accessible')
        key = serialization.load_pem_private_key(path.read_bytes(), password=None)
        if not isinstance(key, Ed25519PrivateKey):
            raise ValueError('Persistent signing key must be Ed25519')
        keys[role] = key
    public_bytes = [key.public_key().public_bytes(serialization.Encoding.Raw, serialization.PublicFormat.Raw)
                    for key in keys.values()]
    if len(set(public_bytes)) != len(PUBLIC_KEYS):
        raise ValueError('Manifest, borrower, and lender must have distinct demo keys')
    return keys


def build_artifacts(job_dir: Path, sources: list[dict], invoices: list, audit: dict,
                    risk: dict, loan_terms: dict, mode: str,
                    template_path: Path | None = None) -> list[str]:
    """Stage, hash, attest, and atomically publish one immutable demo bundle.

    Returns artifact basenames, not private key paths. Existing artifacts are
    never overwritten; failed staging directories are retained for diagnosis.
    Manifest JSON and its signature cannot hash themselves (circular), so the
    signed outputs list covers all other artifacts. Public keys embedded here
    prove self-contained demo integrity, not externally trusted identities.
    Every supplied source name/hash is retained, including extraction, proposal,
    audit, risk, and original upload files. The caller supplies their complete
    inventory; no evidence files or source hashes are fabricated here.
    """
    if not isinstance(audit, dict) or not isinstance(risk, dict):
        raise TypeError('audit and risk must be dictionaries')
    if not isinstance(mode, str) or not mode.strip():
        raise ValueError('A nonempty demo workflow mode is required')
    if set(re.split(r'[^a-z]+', mode.lower())) & {'real', 'live', 'production'}:
        raise ValueError('Only demo workflows are supported; real/live funding is forbidden')
    if not isinstance(sources, list):
        raise TypeError('sources must be a list')
    declared_sources, names = [], set()
    for source in sources:
        if not isinstance(source, dict):
            raise TypeError('Each source must be a dictionary with name and sha256')
        name = _relative_name(source['name'])
        if name in names:
            raise ValueError(f'Duplicate source name: {name}')
        names.add(name)
        declared_sources.append({'name': name, 'sha256': _hash_string(source['sha256'])})
    if template_path is not None:
        template_path = Path(template_path)
        template_hash = _sha256(template_path)
        template_name = _relative_name(template_path.name)
        existing = next((entry for entry in declared_sources if entry['name'] == template_name), None)
        if existing is not None and existing['sha256'] != template_hash:
            raise ValueError('Supplied template hash disagrees with source manifest')
        if existing is None:
            declared_sources.append({'name': template_name, 'sha256': template_hash})
    declared_sources.sort(key=lambda entry: entry['name'])
    model = _ledger_model(invoices, audit.get('bank_rows', []), loan_terms)
    evidence_input = {'sources': declared_sources, 'loan_terms': _json_safe(loan_terms)}
    evidence_hash = hashlib.sha256(canonical_json(evidence_input)).hexdigest()
    # Validate all evidence before creating any staging output.
    canonical_json({'risk': risk, 'audit': audit, 'model': model})
    job_dir = Path(job_dir).resolve()
    artifacts = job_dir / 'artifacts'
    if artifacts.exists() or artifacts.is_symlink():
        raise FileExistsError('Artifacts already exist; use a new job directory, not an overwrite')
    staging_root = job_dir / 'staging'
    if staging_root.is_symlink():
        raise ValueError('Staging directory must not be a symlink')
    staging_root.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix='bundle-', dir=staging_root))
    key_dir = job_dir.parent.parent / 'keys'
    if key_dir == job_dir or job_dir in key_dir.parents or artifacts in key_dir.parents:
        raise ValueError('Private signing keys must live outside job outputs')
    keys = _signing_keys(key_dir)
    serialization, _ = _crypto_dependencies()
    pdf_result = compile_tnc_agreement(loan_terms, str(template_path) if template_path else None,
                                       str(stage / PDF_NAME), evidence_hash)
    build_master_ledger(invoices, audit.get('bank_rows', []), loan_terms, str(stage / LEDGER_NAME))
    if template_path is not None and _sha256(template_path) != template_hash:
        raise ValueError('Template changed during artifact generation; staging retained')
    pdf_bytes = (stage / PDF_NAME).read_bytes()
    for role, name in PUBLIC_KEYS.items():
        (stage / name).write_bytes(keys[role].public_key().public_bytes(
            serialization.Encoding.PEM, serialization.PublicFormat.SubjectPublicKeyInfo))
    for role, name in ROLE_SIGNATURES.items():
        message = ATTESTATION_DOMAIN + role.encode('ascii') + b'\x00' + pdf_bytes
        signature = keys[role].sign(message)
        keys[role].public_key().verify(signature, message)
        (stage / name).write_bytes(signature)
    hashed_names = [LEDGER_NAME, PDF_NAME, *PUBLIC_KEYS.values(), *ROLE_SIGNATURES.values()]
    output_hashes = [{'name': name, 'sha256': _sha256(stage / name)} for name in sorted(hashed_names)]
    if next(entry['sha256'] for entry in output_hashes if entry['name'] == PDF_NAME) != pdf_result['sha256']:
        raise ValueError('PDF changed after compilation')
    manifest = {
        'schema_version': 1, 'mode': mode, 'is_demo': True, 'legal_status': LEGAL_STATUS,
        'real_funds': False, 'sources': declared_sources, 'outputs': output_hashes,
        'source_hash_status': 'declared; template checked if supplied; verifier can check source directory',
        'policy': dict(POLICY),
        'policy_gates': {
            'computed': {'demo_only': True, 'real_funds_permitted': False,
                         'legal_countersignature_claim': False, 'principal_within_cap': True,
                         'principal_within_invoice_support': True,
                         'decimal_allocations_conserve_principal': model['allocation_sum'] == model['loan']['principal']},
            'supplied': {'audit': _json_safe(audit.get('policy_gates', audit.get('gates'))),
                         'risk': _json_safe(risk.get('policy_gates', risk.get('gates'))),
                         'loan_terms': _json_safe(loan_terms.get('policy_gates'))},
            'values': {'principal': model['loan']['principal'], 'loan_cap': model['loan']['loan_cap'],
                       'invoice_support': model['invoice_sum'], 'unmatched_bank_inflow': model['unmatched_bank_inflow']},
        },
        'risk': _json_safe(risk), 'audit_evidence_summary': _json_safe(audit), 'decimal_audit': model,
        'evidence_hash': evidence_hash, 'evidence_hash_input': evidence_input,
        'signatures': {
            'algorithm': 'Ed25519', 'manifest': {'signature': MANIFEST_SIG, 'public_key': PUBLIC_KEYS['manifest'],
                                               'message': 'canonical UTF-8 JSON; sort_keys; separators comma/colon; ensure_ascii=false'},
            'pdf': {role: {'signature': ROLE_SIGNATURES[role], 'public_key': PUBLIC_KEYS[role],
                           'artifact': PDF_NAME, 'message': 'OUT_IN_DEMO_PDF_ATTESTATION_V1 NUL role NUL final PDF bytes'}
                    for role in ROLE_SIGNATURES},
            'trust': 'Demo public keys are self-contained; no identity or legal authority is asserted',
        },
        'hash_exclusions': [MANIFEST_NAME, MANIFEST_SIG],
    }
    manifest_bytes = canonical_json(manifest)
    signature = keys['manifest'].sign(manifest_bytes)
    keys['manifest'].public_key().verify(signature, manifest_bytes)
    (stage / MANIFEST_NAME).write_bytes(manifest_bytes)
    (stage / MANIFEST_SIG).write_bytes(signature)
    # Every payload, including signatures and public keys, is complete before
    # the directory becomes visible. No broad deletes or partial publication.
    if artifacts.exists() or artifacts.is_symlink():
        raise FileExistsError('Artifacts appeared during generation; staging retained')
    os.rename(stage, artifacts)
    return [LEDGER_NAME, PDF_NAME, MANIFEST_NAME, MANIFEST_SIG,
            *ROLE_SIGNATURES.values(), *PUBLIC_KEYS.values()]
