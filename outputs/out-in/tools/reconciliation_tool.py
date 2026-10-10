"""Exact reconciliation. Decimal, not binary floats; no tolerance or hidden rounding.

The supplied legacy Lambda is a DynamoDB report exporter with fabricated audit
rows, not a parity verifier. It is retained under reference/, never executed here.
Statements require transaction_id, invoice_id, currency, inflow, outflow, date.
"""
from __future__ import annotations
import csv
import json
import sys
from collections import defaultdict
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path


class EvidenceError(ValueError):
    pass


class ParityViolation(EvidenceError):
    def __init__(self, report: dict):
        super().__init__("Exact parity failed; downstream execution prohibited")
        self.report = report


def dec(value, label="amount") -> Decimal:
    if isinstance(value, (bool, float)) or value is None:
        raise EvidenceError(f"{label} must be a decimal string or integer, not a float")
    try:
        text = str(value).strip()
        if len(text) > 40:
            raise EvidenceError(f"{label} exceeds numeric bounds")
        result = Decimal(text)
    except InvalidOperation as exc:
        raise EvidenceError(f"Invalid {label}") from exc
    if not result.is_finite() or abs(result) > Decimal("1e15") or result.as_tuple().exponent < -12:
        raise EvidenceError(f"Invalid or out-of-range {label}")
    return result


def money(value, label="amount") -> Decimal:
    result = dec(value, label)
    if result != result.quantize(Decimal("0.01")):
        raise EvidenceError(f"{label} has sub-cent precision; cannot silently round")
    return result


def read_statement(statement_path: str) -> list[dict]:
    path = Path(statement_path)
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if "BankLedger" not in book.sheetnames:
                raise EvidenceError("Statement requires a BankLedger worksheet")
            sheet = book["BankLedger"]
            if sheet.max_row > 10001 or sheet.max_column > 32:
                raise EvidenceError("Statement exceeds row/column limits")
            it = sheet.iter_rows(values_only=False)
            header = [str(c.value or "").strip().lower() for c in next(it)]
            if len(set(header)) != len(header):
                raise EvidenceError("Duplicate statement headers")
            rows = []
            for cells in it:
                if all(c.value is None for c in cells):
                    continue
                if any(c.data_type == "f" for c in cells):
                    raise EvidenceError("Statement formulas are untrusted; supply literal bank values")
                row = dict(zip(header, [c.value for c in cells]))
                # Excel numbers are display decimals. Preserve shortest decimal text;
                # do not use Decimal(binary_float) or a tolerance.
                for key in ("inflow", "outflow"):
                    if isinstance(row.get(key), float):
                        row[key] = str(row[key])
                rows.append(row)
        finally:
            book.close()
    else:
        raise EvidenceError("Statement must be CSV or XLSX")
    if not rows or len(rows) > 10000:
        raise EvidenceError("Statement needs 1..10000 rows")
    required = {"transaction_id", "invoice_id", "currency", "inflow", "outflow", "date"}
    normalized, seen = [], set()
    for row in rows:
        if not required.issubset(row):
            raise EvidenceError("Missing bank statement columns: " + ", ".join(sorted(required - row.keys())))
        tid = str(row["transaction_id"] or "").strip()
        if not tid or tid in seen:
            raise EvidenceError("Missing or duplicate transaction_id")
        seen.add(tid)
        if row["currency"] != "MYR":
            raise EvidenceError("Mixed/foreign currency requires externally verified FX evidence; no implicit conversion")
        inflow, outflow = money(row["inflow"], "inflow"), money(row["outflow"], "outflow")
        if inflow < 0 or outflow < 0 or (inflow > 0 and outflow > 0):
            raise EvidenceError("Bank inflow/outflow direction is invalid")
        normalized.append({**row, "transaction_id": tid,
                           "invoice_id": str(row["invoice_id"] or "").strip(),
                           "inflow": str(inflow), "outflow": str(outflow), "date": str(row["date"])})
    return normalized


def run_excel_reconciliation(invoice_data: list, statement_path: str,
                             delivery_data: list | None = None) -> dict:
    if not invoice_data or len(invoice_data) > 500:
        raise EvidenceError("Dossier requires 1..500 invoices")
    rows = read_statement(statement_path)
    errors, ids, invoice_totals = [], set(), {}
    deltas = []
    with localcontext() as ctx:
        ctx.prec = 80  # bounded inputs guarantee exact products and sums
        for invoice in invoice_data:
            iid = str(invoice.get("invoice_id", ""))
            if not iid or iid in ids:
                raise EvidenceError("Missing or duplicate invoice_id")
            ids.add(iid)
            if invoice.get("currency") != "MYR":
                raise EvidenceError("Only verified MYR invoices are supported; FX gate blocks other currencies")
            items = invoice.get("line_items", [])
            if not items or len(items) > 5000:
                raise EvidenceError("Invoice needs 1..5000 line items")
            products = []
            for line in items:
                qty, price = dec(line.get("quantity"), "quantity"), dec(line.get("unit_price"), "unit_price")
                if qty <= 0 or price < 0 or not str(line.get("sku", "")):
                    raise EvidenceError("Line requires a SKU, positive quantity and nonnegative price")
                products.append(qty * price)
            subtotal = money(invoice.get("subtotal"), "subtotal")
            total = money(invoice.get("total"), "total")
            tax = money(invoice.get("tax", "0"), "tax")
            freight = money(invoice.get("freight", "0"), "freight")
            discount = money(invoice.get("discount", "0"), "discount")
            if min(subtotal, tax, freight, discount) < 0 or total <= 0:
                raise EvidenceError("Invoice totals/adjustments have invalid signs")
            for gate, delta in [("line_items_to_subtotal", sum(products, Decimal(0)) - subtotal),
                                ("subtotal_adjustments_to_total", subtotal + tax + freight - discount - total)]:
                if delta != 0:
                    deltas.append(abs(delta))
                    errors.append({"invoice_id": iid, "gate": gate, "delta": str(delta),
                                   "source": invoice.get("source"), "regions": invoice.get("regions", [])})
            invoice_totals[iid] = total
        bank_totals = defaultdict(Decimal)
        for row in rows:
            inflow = money(row["inflow"])
            if inflow == 0:
                continue
            if row["invoice_id"] not in ids:
                errors.append({"gate": "unmatched_bank_inflow", "transaction_id": row["transaction_id"],
                               "delta": str(inflow)})
                deltas.append(inflow)
            bank_totals[row["invoice_id"]] += inflow
        for iid, total in invoice_totals.items():
            delta = total - bank_totals[iid]
            if delta != 0:
                errors.append({"gate": "invoice_to_bank", "invoice_id": iid, "delta": str(delta)})
                deltas.append(abs(delta))
        grand_total = sum(invoice_totals.values(), Decimal(0))
        inflow_sum = sum((money(r["inflow"]) for r in rows), Decimal(0))
        ledger_delta = grand_total - inflow_sum
        if ledger_delta:
            deltas.append(abs(ledger_delta))
        if delivery_data is not None:
            delivered = defaultdict(Decimal)
            for docket in delivery_data:
                iid = docket.get("invoice_id")
                if iid not in ids:
                    raise EvidenceError("Delivery docket does not match an invoice")
                for line in docket.get("line_items", []):
                    qty = dec(line.get("quantity"), "delivered quantity")
                    if qty <= 0:
                        raise EvidenceError("Invalid delivery quantity")
                    delivered[(iid, line.get("sku"))] += qty
            invoiced = defaultdict(Decimal)
            for invoice in invoice_data:
                iid = invoice["invoice_id"]
                for line in invoice["line_items"]:
                    invoiced[(iid, line["sku"])] += dec(line["quantity"])
                dockets = [d for d in delivery_data if d.get("invoice_id") == iid]
                if not dockets:
                    errors.append({"gate": "missing_delivery_docket", "invoice_id": iid})
                for field in ("container_seal", "incoterms"):
                    if not invoice.get(field) or any(d.get(field) != invoice[field] for d in dockets):
                        errors.append({"gate": "shipping_identity", "invoice_id": iid, "field": field})
            for key in invoiced.keys() | delivered.keys():
                delta = invoiced[key] - delivered[key]
                if delta:
                    errors.append({"gate": "delivery_quantity", "invoice_id": key[0], "sku": key[1],
                                   "quantity_delta": str(delta)})
        return {"passed": not errors and ledger_delta == 0, "variance": str(max(deltas, default=Decimal(0))),
                "ledger_delta": str(ledger_delta), "invoice_total": str(grand_total),
                "verified_inflow": str(inflow_sum), "errors": errors, "bank_rows": rows,
                "gates": {"exact_money_parity": not deltas, "delivery_verified": (not errors) if delivery_data is not None else None},
                "arithmetic": "bounded Decimal precision=80; equality exact; zero tolerance",
                "statement_sha256": __import__("hashlib").sha256(Path(statement_path).read_bytes()).hexdigest()}


def enforce_parity(report: dict) -> dict:
    # Independent guard: neither LLM text nor a hook can override the scheduler.
    required = {"passed", "variance", "ledger_delta", "invoice_total", "verified_inflow", "errors", "gates", "bank_rows"}
    if not isinstance(report, dict) or not required.issubset(report):
        raise EvidenceError("Incomplete parity report")
    gates = report.get("gates")
    if not isinstance(gates, dict) or gates.get("exact_money_parity") is not True or gates.get("delivery_verified") is False:
        raise ParityViolation(report)
    if not isinstance(report.get("errors"), list) or report["errors"]:
        raise ParityViolation(report)
    if report.get("passed") is not True or dec(report.get("variance")) != 0 or dec(report.get("ledger_delta")) != 0 or dec(report["invoice_total"]) != dec(report["verified_inflow"]):
        raise ParityViolation(report)
    return report


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: python -m tools.reconciliation_tool invoices.json statement.xlsx")
    result = run_excel_reconciliation(json.loads(Path(sys.argv[1]).read_text(encoding="utf-8")), sys.argv[2])
    print(json.dumps(result, indent=2))
    raise SystemExit(0 if result["passed"] else 2)
